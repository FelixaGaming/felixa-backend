"""
FELIXA AUTOMATED BACKEND
========================
Fully automated community health analysis system.

Flow:
1. Customer fills form on website
2. Pays via Stripe
3. Stripe webhook triggers analysis
4. Report emailed to customer + play@felixagaming.com

Deploy to: Railway (free tier)
"""

import os
import json
import asyncio
import requests
from datetime import datetime
from typing import Optional, List, Dict, Any
from concurrent.futures import ThreadPoolExecutor
import base64
from io import BytesIO

from fastapi import FastAPI, HTTPException, Request, BackgroundTasks
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from pydantic import BaseModel
import stripe
import resend

from openai import OpenAI

# ============================================
# CONFIGURATION
# ============================================

# API Keys (set in Railway environment variables)
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
STRIPE_SECRET_KEY = os.getenv("STRIPE_SECRET_KEY")
STRIPE_WEBHOOK_SECRET = os.getenv("STRIPE_WEBHOOK_SECRET")
RESEND_API_KEY = os.getenv("RESEND_API_KEY")
YOUTUBE_API_KEY = os.getenv("YOUTUBE_API_KEY", "")
REDDIT_CLIENT_ID = os.getenv("REDDIT_CLIENT_ID", "")
REDDIT_CLIENT_SECRET = os.getenv("REDDIT_CLIENT_SECRET", "")

# Email settings
ADMIN_EMAIL = "play@felixagaming.com"
FROM_EMAIL = "Felixa <reports@felixagaming.com>"  # You'll verify this domain in Resend

# Initialize clients
openai_client = OpenAI(api_key=OPENAI_API_KEY) if OPENAI_API_KEY else None
stripe.api_key = STRIPE_SECRET_KEY
resend.api_key = RESEND_API_KEY

thread_pool = ThreadPoolExecutor(max_workers=5)

# ============================================
# FASTAPI APP
# ============================================

app = FastAPI(title="Felixa Automated Backend", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ============================================
# ATTRIBUTES & CATEGORIES (matches Excel)
# ============================================

ATTRIBUTE_WEIGHTS = {
    "Polite": 3, "Empathetic": 3, "Encouraging": 2, "Funny": 1,
    "Ignorant": -1, "Ego-centric": -1, "Neurotic": -1,
    "Sarcasm": -2, "Agitated": -2,
    "Aggressive": -3, "Judgmental": -3, "Disrespectful": -3, "Rude": -3,
}

ALL_ATTRIBUTES = list(ATTRIBUTE_WEIGHTS.keys())
CATEGORIES = ["Profanity", "Hate_Speech", "Violence", "Spam"]

# ============================================
# GPT ANALYSIS PROMPT
# ============================================

ANALYSIS_PROMPT = """Analyze the following messages for behavioral attributes and content categories.

POSITIVE ATTRIBUTES:
• Polite: Respectful and courteous • Funny: Light-hearted • Empathetic: Shows care • Encouraging: Supportive

NEGATIVE ATTRIBUTES:
• Low (-1): Ignorant, Ego-centric, Neurotic
• Medium (-2): Sarcasm, Agitated  
• High (-3): Aggressive, Judgmental, Disrespectful, Rude

CONTENT CATEGORIES:
1. Profanity: Swear words
2. Hate_Speech: Discrimination
3. Violence: Threats
4. Spam: Promotional/scam content

Return ONLY valid JSON:
{
    "comments": [
        {
            "comment_index": 0,
            "attributes": {"Polite": 0, "Funny": 0, "Empathetic": 0, "Encouraging": 0, "Ignorant": 0, "Ego-centric": 0, "Neurotic": 0, "Sarcasm": 0, "Agitated": 0, "Aggressive": 0, "Judgmental": 0, "Disrespectful": 0, "Rude": 0},
            "categories": {"Profanity": {"present": 0, "words": []}, "Hate_Speech": {"present": 0, "words": []}, "Violence": {"present": 0, "words": []}, "Spam": {"present": 0, "words": []}},
            "sentiment": "neutral"
        }
    ]
}

MESSAGES:
"""

# ============================================
# SCRAPING FUNCTIONS
# ============================================

def scrape_youtube(url: str, limit: int = 500) -> List[Dict]:
    """Scrape YouTube comments"""
    if not YOUTUBE_API_KEY:
        return []
    
    try:
        video_id = None
        if "v=" in url:
            video_id = url.split("v=")[1].split("&")[0]
        elif "youtu.be/" in url:
            video_id = url.split("youtu.be/")[1].split("?")[0]
        
        if not video_id:
            return []
        
        comments = []
        params = {
            "part": "snippet",
            "videoId": video_id,
            "key": YOUTUBE_API_KEY,
            "maxResults": 100,
            "textFormat": "plainText",
        }
        
        while len(comments) < limit:
            r = requests.get("https://www.googleapis.com/youtube/v3/commentThreads", params=params)
            if r.status_code != 200:
                break
            
            data = r.json()
            for item in data.get("items", []):
                text = item["snippet"]["topLevelComment"]["snippet"].get("textDisplay", "")
                if text:
                    comments.append({"text": text})
            
            if "nextPageToken" in data and len(comments) < limit:
                params["pageToken"] = data["nextPageToken"]
            else:
                break
        
        return comments[:limit]
    except Exception as e:
        print(f"YouTube scrape error: {e}")
        return []


def scrape_reddit(url: str, limit: int = 500) -> List[Dict]:
    """Scrape Reddit comments"""
    if not REDDIT_CLIENT_ID or not REDDIT_CLIENT_SECRET:
        return []
    
    try:
        import praw
        
        reddit = praw.Reddit(
            client_id=REDDIT_CLIENT_ID,
            client_secret=REDDIT_CLIENT_SECRET,
            user_agent="Felixa/1.0"
        )
        
        submission = reddit.submission(url=url)
        submission.comments.replace_more(limit=0)
        
        comments = []
        for comment in submission.comments.list()[:limit]:
            if hasattr(comment, 'body') and comment.body:
                comments.append({"text": comment.body})
        
        return comments
    except Exception as e:
        print(f"Reddit scrape error: {e}")
        return []


def scrape_platform(platform: str, url: str, limit: int = 500) -> List[Dict]:
    """Route to correct scraper"""
    platform = platform.lower()
    
    if platform == "youtube":
        return scrape_youtube(url, limit)
    elif platform == "reddit":
        return scrape_reddit(url, limit)
    elif platform == "twitch":
        # Twitch requires more complex setup
        return []
    elif platform == "discord":
        # Discord handled separately via bot
        return []
    else:
        return []


# ============================================
# GPT ANALYSIS
# ============================================

async def analyze_comments(comments: List[Dict]) -> Dict:
    """Analyze comments with GPT-4 Mini"""
    if not openai_client or not comments:
        return {"comments": []}
    
    all_results = []
    batch_size = 15
    
    for i in range(0, len(comments), batch_size):
        batch = comments[i:i + batch_size]
        batch_text = "\n".join([f'Message {j}: "{c["text"][:400]}"' for j, c in enumerate(batch)])
        
        try:
            loop = asyncio.get_event_loop()
            response = await loop.run_in_executor(
                thread_pool,
                lambda: openai_client.chat.completions.create(
                    model="gpt-4o-mini",
                    messages=[
                        {"role": "system", "content": "Analyze text for behavioral attributes. Return only valid JSON."},
                        {"role": "user", "content": ANALYSIS_PROMPT + batch_text}
                    ],
                    temperature=0.0,
                    max_tokens=4000
                )
            )
            
            content = response.choices[0].message.content.strip()
            if content.startswith("```"):
                content = content.split("```")[1]
                if content.startswith("json"):
                    content = content[4:]
            
            result = json.loads(content)
            for j, item in enumerate(result.get("comments", [])):
                item["comment_index"] = i + j
                all_results.append(item)
                
        except Exception as e:
            print(f"Analysis error: {e}")
        
        await asyncio.sleep(0.3)
    
    return {"comments": all_results}


def calculate_health_score(results: Dict) -> Dict:
    """Calculate health score from results"""
    comments = results.get("comments", [])
    total = len(comments)
    
    if total == 0:
        return {"health_score": 50, "health_status": "No Data", "summary": {}}
    
    attr_counts = {a: 0 for a in ALL_ATTRIBUTES}
    cat_counts = {c: 0 for c in CATEGORIES}
    sent_counts = {"positive": 0, "neutral": 0, "negative": 0}
    
    for c in comments:
        for attr, val in c.get("attributes", {}).items():
            if val == 1 and attr in attr_counts:
                attr_counts[attr] += 1
        
        for cat in CATEGORIES:
            cat_data = c.get("categories", {}).get(cat, {})
            if isinstance(cat_data, dict) and cat_data.get("present", 0) == 1:
                cat_counts[cat] += 1
        
        sent = c.get("sentiment", "neutral").lower()
        if sent in sent_counts:
            sent_counts[sent] += 1
    
    pos_weight = sum(attr_counts[a] * ATTRIBUTE_WEIGHTS[a] for a in ALL_ATTRIBUTES if ATTRIBUTE_WEIGHTS[a] > 0)
    neg_weight = abs(sum(attr_counts[a] * ATTRIBUTE_WEIGHTS[a] for a in ALL_ATTRIBUTES if ATTRIBUTE_WEIGHTS[a] < 0))
    
    max_pos = total * 9
    score = min(100, max(0, int(50 + (pos_weight - neg_weight) / max(1, max_pos) * 50)))
    
    status = "Excellent" if score >= 80 else "Good" if score >= 60 else "Fair" if score >= 40 else "Needs Attention"
    
    return {
        "health_score": score,
        "health_status": status,
        "summary": {
            "total": total,
            "positive": sent_counts["positive"],
            "negative": sent_counts["negative"],
            "attributes": attr_counts,
            "categories": cat_counts
        }
    }


# ============================================
# EXCEL GENERATION
# ============================================

def generate_excel(results: Dict, comments: List[Dict], platform: str, source: str) -> bytes:
    """Generate Excel report and return as bytes"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    wb = Workbook()
    
    # Colors
    PURPLE = "6B4C9A"
    GREEN = "28A745"
    RED = "DC3545"
    WHITE = "FFFFFF"
    GRAY = "6C757D"
    
    border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )
    
    health = calculate_health_score(results)
    
    # Dashboard Sheet
    ws = wb.active
    ws.title = "Dashboard"
    
    ws.merge_cells('A1:F1')
    ws['A1'] = "FELIXA COMMUNITY HEALTH REPORT"
    ws['A1'].font = Font(bold=True, size=20, color=WHITE)
    ws['A1'].fill = PatternFill('solid', fgColor=PURPLE)
    ws['A1'].alignment = Alignment(horizontal='center')
    
    ws['A3'] = "Platform:"
    ws['B3'] = platform.upper()
    ws['A4'] = "Source:"
    ws['B4'] = source[:50]
    ws['A5'] = "Date:"
    ws['B5'] = datetime.now().strftime("%Y-%m-%d %H:%M")
    ws['A6'] = "Messages Analyzed:"
    ws['B6'] = health["summary"].get("total", 0)
    
    ws['A8'] = "HEALTH SCORE"
    ws['A8'].font = Font(bold=True, size=14)
    ws['A9'] = health["health_score"]
    ws['A9'].font = Font(bold=True, size=36, color=GREEN if health["health_score"] >= 60 else RED)
    ws['B9'] = health["health_status"]
    ws['B9'].font = Font(bold=True, size=16)
    
    ws['A11'] = "Positive Messages:"
    ws['B11'] = health["summary"].get("positive", 0)
    ws['A12'] = "Negative Messages:"
    ws['B12'] = health["summary"].get("negative", 0)
    
    # Attributes Sheet
    ws2 = wb.create_sheet("Attributes")
    ws2['A1'] = "Attribute"
    ws2['B1'] = "Count"
    ws2['C1'] = "Weight"
    
    for i, (attr, weight) in enumerate(ATTRIBUTE_WEIGHTS.items(), start=2):
        ws2[f'A{i}'] = attr
        ws2[f'B{i}'] = health["summary"].get("attributes", {}).get(attr, 0)
        ws2[f'C{i}'] = weight
    
    # Categories Sheet
    ws3 = wb.create_sheet("Categories")
    ws3['A1'] = "Category"
    ws3['B1'] = "Count"
    
    for i, cat in enumerate(CATEGORIES, start=2):
        ws3[f'A{i}'] = cat
        ws3[f'B{i}'] = health["summary"].get("categories", {}).get(cat, 0)
    
    # Flagged Comments Sheet
    ws4 = wb.create_sheet("Flagged Comments")
    ws4['A1'] = "#"
    ws4['B1'] = "Comment"
    ws4['C1'] = "Flags"
    ws4['D1'] = "Severity"
    
    flagged = []
    backend_comments = results.get("comments", [])
    
    high_attrs = ["Aggressive", "Judgmental", "Disrespectful", "Rude"]
    high_cats = ["Hate_Speech", "Violence"]
    
    for i, (orig, analyzed) in enumerate(zip(comments[:len(backend_comments)], backend_comments)):
        attrs = analyzed.get("attributes", {})
        cats = analyzed.get("categories", {})
        
        flags = []
        severity = "Medium"
        
        for attr in high_attrs:
            if attrs.get(attr, 0) == 1:
                flags.append(attr)
                severity = "High"
        
        for cat in high_cats:
            cat_data = cats.get(cat, {})
            if isinstance(cat_data, dict) and cat_data.get("present", 0) == 1:
                flags.append(cat.replace("_", " "))
                severity = "High"
        
        if not flags:
            prof = cats.get("Profanity", {})
            if isinstance(prof, dict) and prof.get("present", 0) == 1:
                flags.append("Profanity")
        
        if flags:
            flagged.append({
                "text": orig.get("text", "")[:150],
                "flags": ", ".join(flags),
                "severity": severity
            })
    
    flagged.sort(key=lambda x: 0 if x["severity"] == "High" else 1)
    
    for i, f in enumerate(flagged[:20], start=2):
        ws4[f'A{i}'] = i - 1
        ws4[f'B{i}'] = f["text"]
        ws4[f'C{i}'] = f["flags"]
        ws4[f'D{i}'] = f["severity"]
    
    ws4.column_dimensions['B'].width = 60
    ws4.column_dimensions['C'].width = 25
    
    # Save to bytes
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# ============================================
# EMAIL SENDING
# ============================================

def send_report_email(
    customer_email: str,
    platform: str,
    source: str,
    health_score: int,
    excel_bytes: bytes
):
    """Send report to customer and admin"""
    
    excel_base64 = base64.b64encode(excel_bytes).decode()
    
    html_content = f"""
    <div style="font-family: Arial, sans-serif; max-width: 600px; margin: 0 auto;">
        <div style="background: #6B4C9A; padding: 20px; text-align: center;">
            <h1 style="color: white; margin: 0;">🩺 Felixa Health Report</h1>
        </div>
        
        <div style="padding: 30px; background: #f8f9fa;">
            <h2 style="color: #2D1B4E;">Your Community Health Report is Ready!</h2>
            
            <div style="background: white; padding: 20px; border-radius: 10px; margin: 20px 0;">
                <p><strong>Platform:</strong> {platform.upper()}</p>
                <p><strong>Source:</strong> {source[:100]}</p>
                <p><strong>Health Score:</strong> <span style="font-size: 24px; color: {'#28A745' if health_score >= 60 else '#DC3545'};">{health_score}/100</span></p>
            </div>
            
            <p>Your detailed Excel report is attached to this email. It includes:</p>
            <ul>
                <li>📊 Dashboard with key metrics</li>
                <li>📈 Behavioral attributes breakdown</li>
                <li>⚠️ Content categories analysis</li>
                <li>🚩 Flagged comments for review</li>
            </ul>
            
            <p style="color: #666; font-size: 14px; margin-top: 30px;">
                Thank you for using Felixa! If you have questions, reply to this email.
            </p>
        </div>
        
        <div style="background: #2D1B4E; padding: 15px; text-align: center;">
            <p style="color: white; margin: 0; font-size: 12px;">
                © 2026 Felixa Gaming | www.felixagaming.com
            </p>
        </div>
    </div>
    """
    
    try:
        # Send to customer
        resend.Emails.send({
            "from": FROM_EMAIL,
            "to": customer_email,
            "subject": f"🩺 Your Felixa Health Report - Score: {health_score}/100",
            "html": html_content,
            "attachments": [
                {
                    "filename": f"felixa_report_{datetime.now().strftime('%Y%m%d')}.xlsx",
                    "content": excel_base64,
                }
            ]
        })
        print(f"✅ Email sent to customer: {customer_email}")
        
        # Send copy to admin
        resend.Emails.send({
            "from": FROM_EMAIL,
            "to": ADMIN_EMAIL,
            "subject": f"[COPY] Felixa Report for {customer_email} - {platform}",
            "html": f"<p>Copy of report sent to {customer_email}</p>" + html_content,
            "attachments": [
                {
                    "filename": f"felixa_report_{datetime.now().strftime('%Y%m%d')}.xlsx",
                    "content": excel_base64,
                }
            ]
        })
        print(f"✅ Copy sent to admin: {ADMIN_EMAIL}")
        
        return True
    except Exception as e:
        print(f"❌ Email error: {e}")
        return False


# ============================================
# MAIN PROCESSING FUNCTION
# ============================================

async def process_order(customer_email: str, platform: str, url: str):
    """Process a paid order"""
    print(f"📦 Processing order: {customer_email} | {platform} | {url}")
    
    # Step 1: Scrape comments
    print("📥 Scraping comments...")
    comments = scrape_platform(platform, url, limit=500)
    
    if not comments:
        print("⚠️ No comments found, using sample data")
        comments = [{"text": "Sample comment for testing"}]
    
    print(f"   Found {len(comments)} comments")
    
    # Step 2: Analyze with GPT
    print("🔍 Analyzing with GPT...")
    results = await analyze_comments(comments)
    print(f"   Analyzed {len(results.get('comments', []))} comments")
    
    # Step 3: Calculate health score
    health = calculate_health_score(results)
    print(f"   Health Score: {health['health_score']}")
    
    # Step 4: Generate Excel
    print("📊 Generating Excel report...")
    excel_bytes = generate_excel(results, comments, platform, url)
    
    # Step 5: Send email
    print("📧 Sending email...")
    send_report_email(
        customer_email=customer_email,
        platform=platform,
        source=url,
        health_score=health["health_score"],
        excel_bytes=excel_bytes
    )
    
    print("✅ Order complete!")
    return {"status": "success", "health_score": health["health_score"]}


# ============================================
# API ENDPOINTS
# ============================================

@app.get("/")
async def root():
    return {
        "service": "Felixa Automated Backend",
        "status": "running",
        "endpoints": ["/webhook/stripe", "/test"]
    }


@app.post("/webhook/stripe")
async def stripe_webhook(request: Request, background_tasks: BackgroundTasks):
    """Handle Stripe webhook after payment"""
    payload = await request.body()
    sig_header = request.headers.get("stripe-signature")
    
    try:
        # Verify webhook signature
        if STRIPE_WEBHOOK_SECRET:
            event = stripe.Webhook.construct_event(
                payload, sig_header, STRIPE_WEBHOOK_SECRET
            )
        else:
            event = json.loads(payload)
        
        # Handle checkout.session.completed
        if event["type"] == "checkout.session.completed":
            session = event["data"]["object"]
            
            # Extract customer data from metadata
            customer_email = session.get("customer_email") or session.get("customer_details", {}).get("email")
            metadata = session.get("metadata", {})
            
            platform = metadata.get("platform", "youtube")
            url = metadata.get("url", "")
            
            print(f"💰 Payment received: {customer_email}")
            
            # Process in background
            background_tasks.add_task(
                process_order,
                customer_email=customer_email,
                platform=platform,
                url=url
            )
            
            return {"status": "processing"}
        
        return {"status": "ignored", "type": event["type"]}
        
    except Exception as e:
        print(f"Webhook error: {e}")
        raise HTTPException(400, str(e))


@app.post("/test")
async def test_analysis(
    email: str = "test@example.com",
    platform: str = "youtube",
    url: str = "https://youtube.com/watch?v=test"
):
    """Test endpoint - manually trigger analysis"""
    result = await process_order(email, platform, url)
    return result


@app.get("/health")
async def health_check():
    return {
        "status": "healthy",
        "openai": bool(OPENAI_API_KEY),
        "stripe": bool(STRIPE_SECRET_KEY),
        "resend": bool(RESEND_API_KEY),
        "youtube": bool(YOUTUBE_API_KEY),
        "reddit": bool(REDDIT_CLIENT_ID)
    }


# ============================================
# RUN SERVER
# ============================================

if __name__ == "__main__":
    import uvicorn
    port = int(os.getenv("PORT", 8000))
    print(f"🚀 Starting Felixa Backend on port {port}")
    uvicorn.run(app, host="0.0.0.0", port=port)
